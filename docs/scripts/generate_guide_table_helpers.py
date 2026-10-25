# -*- coding: utf-8 -*-
"""
Table and Cell helper functions with generous row heights and 100% Meiryo UI
"""

import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- Monotone Color Palette ---
C_TITLE_BG     = "1F1F1F"  # Deep Charcoal Black
C_TITLE_FG     = "FFFFFF"  # White
C_SUBTITLE_BG  = "333333"  # Dark Slate
C_SUBTITLE_FG  = "E0E0E0"  # Off-White
C_SECTION_BG   = "2D2D2D"  # Charcoal Section Banner
C_SECTION_FG   = "FFFFFF"  # White
C_TH_BG        = "424242"  # Slate Gray Table Header
C_TH_FG        = "FFFFFF"  # White
C_SUB_TH_BG    = "555555"  # Medium Gray Subheader
C_SUB_TH_FG    = "FFFFFF"  # White
C_ZEBRA_WHITE  = "FFFFFF"  # Pure White
C_ZEBRA_GRAY   = "F5F5F5"  # Soft Gray
C_CALLOUT_BG   = "ECECEC"  # Soft Platinum Gray
C_CALLOUT_FG   = "1F1F1F"  # Dark Charcoal Text
C_CODE_BG      = "F0F0F0"  # Light Gray Code
C_BORDER_LIGHT = "D0D0D0"  # Light Gray Grid Border
C_BORDER_DARK  = "707070"  # Medium Border for Headers

def font_m(size=9.5, bold=False, italic=False, color="000000"):
    return Font(name="Meiryo UI", size=size, bold=bold, italic=italic, color=color)

def fill_s(color_hex):
    return PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")

def border_t():
    s = Side(style="thin", color=C_BORDER_LIGHT)
    return Border(left=s, right=s, top=s, bottom=s)

def border_h():
    so = Side(style="medium", color=C_BORDER_DARK)
    si = Side(style="thin", color=C_BORDER_LIGHT)
    return Border(left=si, right=si, top=so, bottom=so)

def border_c():
    sl = Side(style="thick", color=C_SECTION_BG)
    so = Side(style="thin", color=C_BORDER_LIGHT)
    return Border(left=sl, right=so, top=so, bottom=so)

def align_c(h="left", v="top", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def get_char_display_width(text):
    w = 0.0
    for ch in str(text):
        if ord(ch) > 0x7F:
            w += 2.0
        else:
            w += 1.0
    return w

def calc_needed_lines(text, col_width):
    if not text:
        return 1
    total_lines = 0
    for paragraph in str(text).split("\n"):
        if not paragraph:
            total_lines += 1
            continue
        p_width = get_char_display_width(paragraph)
        capacity = max(1.0, col_width * 1.45)
        wrapped = math.ceil(p_width / capacity)
        total_lines += max(1, wrapped)
    return total_lines

def auto_set_row_height(ws, row_idx, min_height=32.0, extra_padding=16.0, line_height=18.0):
    max_lines = 1
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_idx, column=c)
        if cell.value is not None:
            col_letter = get_column_letter(c)
            cw = ws.column_dimensions[col_letter].width or 20.0
            lines = calc_needed_lines(cell.value, cw)
            max_lines = max(max_lines, lines)
    calculated_height = max(min_height, max_lines * line_height + extra_padding)
    ws.row_dimensions[row_idx].height = calculated_height

def add_header_banner(ws, title, subtitle):
    ws.merge_cells("A1:F1")
    t = ws.cell(1, 1, title)
    t.font = font_m(size=14, bold=True, color=C_TITLE_FG)
    t.fill = fill_s(C_TITLE_BG)
    t.alignment = align_c(h="left", v="center", wrap=False)
    ws.row_dimensions[1].height = 42.0
    for c in range(1, 7):
        ws.cell(1, c).fill = fill_s(C_TITLE_BG)
        ws.cell(1, c).font = font_m(size=14, bold=True, color=C_TITLE_FG)

    ws.merge_cells("A2:F2")
    s = ws.cell(2, 1, subtitle)
    s.font = font_m(size=9.5, bold=False, color=C_SUBTITLE_FG)
    s.fill = fill_s(C_SUBTITLE_BG)
    s.alignment = align_c(h="left", v="center", wrap=False)
    ws.row_dimensions[2].height = 34.0
    for c in range(1, 7):
        ws.cell(2, c).fill = fill_s(C_SUBTITLE_BG)
        ws.cell(2, c).font = font_m(size=9.5, bold=False, color=C_SUBTITLE_FG)

    ws.row_dimensions[3].height = 10.0

def add_section(ws, r, title):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(r, 1, title)
    c.font = font_m(size=11, bold=True, color=C_SECTION_FG)
    c.fill = fill_s(C_SECTION_BG)
    c.alignment = align_c(h="left", v="center", wrap=False)
    ws.row_dimensions[r].height = 34.0
    for col in range(1, 7):
        ws.cell(r, col).fill = fill_s(C_SECTION_BG)
        ws.cell(r, col).font = font_m(size=11, bold=True, color=C_SECTION_FG)
    return r + 1

def add_callout(ws, r, text):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(r, 1, text)
    c.font = font_m(size=9.5, bold=False, color=C_CALLOUT_FG)
    c.fill = fill_s(C_CALLOUT_BG)
    c.alignment = align_c(h="left", v="top", wrap=True)
    b = border_c()
    for col in range(1, 7):
        ws.cell(r, col).fill = fill_s(C_CALLOUT_BG)
        ws.cell(r, col).font = font_m(size=9.5, bold=False, color=C_CALLOUT_FG)
        ws.cell(r, col).border = b
    total_w = sum([ws.column_dimensions[get_column_letter(col)].width or 20 for col in range(1, 7)])
    lines = calc_needed_lines(text, total_w)
    ws.row_dimensions[r].height = max(32.0, lines * 19.0 + 18.0)
    return r + 1

def add_table_header(ws, r, headers, bg_color=C_TH_BG, fg_color=C_TH_FG):
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(r, c_idx, h)
        cell.font = font_m(size=10, bold=True, color=fg_color)
        cell.fill = fill_s(bg_color)
        cell.alignment = align_c(h="center", v="center", wrap=True)
        cell.border = border_h()
    auto_set_row_height(ws, r, min_height=34.0, extra_padding=16.0, line_height=18.0)
    return r + 1

def add_table_row(ws, r, values, is_zebra=False, align_list=None, bold_cols=None):
    fill_color = C_ZEBRA_GRAY if is_zebra else C_ZEBRA_WHITE
    for c_idx, val in enumerate(values, start=1):
        cell = ws.cell(r, c_idx, val)
        is_bold = (bold_cols is not None and c_idx in bold_cols)
        cell.font = font_m(size=9.5, bold=is_bold, color="000000")
        cell.fill = fill_s(fill_color)
        cell.border = border_t()
        h_align = align_list[c_idx - 1] if (align_list and c_idx - 1 < len(align_list)) else "left"
        cell.alignment = align_c(h=h_align, v="top", wrap=True)
    auto_set_row_height(ws, r, min_height=32.0, extra_padding=16.0, line_height=18.0)
    return r + 1
